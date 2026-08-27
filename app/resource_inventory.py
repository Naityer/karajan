"""Auditable XLSX inventory snapshots for the Inspiration catalog."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from threading import RLock
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.resource_validity import decorate_catalog_resources, load_policy

_SNAPSHOT_TTL = timedelta(minutes=10)
_LOCK = RLock()


@dataclass(frozen=True)
class InventorySnapshot:
    id: str
    filename: str
    created_at: datetime
    rows: list[dict]
    policy: dict
    content: bytes


_SNAPSHOTS: dict[str, InventorySnapshot] = {}


INVENTORY_COLUMNS = [
    ("name", "Nombre"), ("kind", "Tipo"), ("parent", "Padre / colección"),
    ("source", "GitHub / fuente"), ("stars", "Estrellas"), ("category", "Categoría"),
    ("runtimes", "Compatibilidad"), ("group", "Grupo"), ("installation", "Instalación"),
    ("validity", "Estado operativo"), ("reasons", "Motivos"), ("checked_at", "Última comprobación"),
]


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _category(item) -> str:
    metadata = item.metadata or {}
    taxonomy = metadata.get("taxonomy") or {}
    intelligence = metadata.get("intelligence") or {}
    return str(taxonomy.get("category") or intelligence.get("primary_taxonomy") or metadata.get("category") or "")


def _inventory_rows(store) -> tuple[list[dict], dict]:
    policy = load_policy(store)
    items = decorate_catalog_resources(store.list_all_resources(include_inactive=True), policy)
    by_id = {item.id: item for item in items}
    group_by_resource: dict[str, str] = {}
    for group in store.list_groups():
        for resource_id in group.resource_ids:
            group_by_resource[resource_id] = group.name
    rows = []
    for item in sorted(items, key=lambda value: (value.name.casefold(), value.id)):
        validity = item.metadata.get("validity") or {}
        parent = by_id.get(item.parent_resource_id) if item.parent_resource_id else None
        rows.append({
            "id": item.id,
            "name": item.name,
            "kind": item.kind.value,
            "parent": parent.name if parent else "",
            "source": validity.get("source_url") or "",
            "stars": validity.get("github_stars") if validity.get("github_repository") else "No aplica",
            "category": _category(item),
            "runtimes": ", ".join("ChatGPT/Codex" if value == "codex" else value.title() for value in item.compatible_runtimes),
            "group": group_by_resource.get(item.id, "Sin clasificar"),
            "installation": item.status,
            "validity": validity.get("state") or "pending_review",
            "reasons": " ".join(validity.get("reasons") or []),
            "checked_at": item.metadata.get("source_analyzed_at") or validity.get("evaluated_at") or "",
        })
    return rows, policy.model_dump(mode="json")


def _workbook_bytes(rows: list[dict], policy: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventario"
    headers = [label for _, label in INVENTORY_COLUMNS]
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(key) for key, _ in INVENTORY_COLUMNS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if rows:
        table = Table(displayName="InventarioKarajan", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        sheet.add_table(table)
    header_fill = PatternFill("solid", fgColor="2B3038")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    widths = [34, 14, 30, 54, 12, 24, 24, 26, 16, 20, 58, 24]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    source_column = next(index for index, (key, _) in enumerate(INVENTORY_COLUMNS, start=1) if key == "source")
    for row_index in range(2, len(rows) + 2):
        cell = sheet.cell(row=row_index, column=source_column)
        if cell.value:
            cell.hyperlink = str(cell.value)
            cell.style = "Hyperlink"
        sheet.cell(row=row_index, column=11).alignment = Alignment(wrap_text=True, vertical="top")
    policy_sheet = workbook.create_sheet("Política")
    policy_sheet.append(["Criterio", "Valor"])
    labels = {
        "version": "Versión", "stars_enabled": "Comprobar estrellas (solo GitHub)", "min_github_stars": "Umbral de estrellas GitHub (> valor)",
        "require_source_url": "Exigir fuente HTTPS", "require_detection_success": "Exigir detección",
        "required_runtimes": "Runtimes requeridos", "require_web_reference": "Exigir referencia web", "updated_at": "Actualizada",
    }
    for key, label in labels.items():
        value = policy.get(key)
        policy_sheet.append([label, ", ".join(value) if isinstance(value, list) else value])
    for cell in policy_sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    policy_sheet.freeze_panes = "A2"
    policy_sheet.column_dimensions["A"].width = 34
    policy_sheet.column_dimensions["B"].width = 42
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_inventory_snapshot(store) -> InventorySnapshot:
    rows, policy = _inventory_rows(store)
    now = _utcnow()
    snapshot = InventorySnapshot(
        id=f"iexp_{uuid4().hex[:12]}",
        filename=f"karajan-inspiration-inventory-{now.date().isoformat()}.xlsx",
        created_at=now,
        rows=rows,
        policy=policy,
        content=_workbook_bytes(rows, policy),
    )
    with _LOCK:
        expired = [key for key, value in _SNAPSHOTS.items() if now - value.created_at > _SNAPSHOT_TTL]
        for key in expired:
            _SNAPSHOTS.pop(key, None)
        _SNAPSHOTS[snapshot.id] = snapshot
    return snapshot


def get_inventory_snapshot(snapshot_id: str) -> InventorySnapshot:
    with _LOCK:
        snapshot = _SNAPSHOTS.get(snapshot_id)
        if snapshot is None or _utcnow() - snapshot.created_at > _SNAPSHOT_TTL:
            _SNAPSHOTS.pop(snapshot_id, None)
            raise KeyError(snapshot_id)
        return snapshot


__all__ = ["INVENTORY_COLUMNS", "create_inventory_snapshot", "get_inventory_snapshot"]
