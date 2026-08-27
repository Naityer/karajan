from __future__ import annotations

from io import BytesIO

import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import main
from app.models import ResourceItem, ResourceKind
from app.nested_resources import ClaudeSkillsCollectionAdapter
from app.resource_library import ResourceStore
from app.resource_validity import decorate_catalog_resources, load_policy


def _client(karajan_install, monkeypatch, tmp_path) -> tuple[TestClient, ResourceStore]:
    resources = ResourceStore(tmp_path / "validity.db")
    karajan_install.runtime.resources = resources
    monkeypatch.setattr(main, "resource_store", resources)
    return karajan_install.client, resources


def _resource(*, slug: str, stars: int, skill: bool = False) -> ResourceItem:
    return ResourceItem(
        slug=slug,
        name=slug.replace("-", " ").title(),
        kind=ResourceKind.SKILL if skill else ResourceKind.REFERENCE,
        source="github",
        repo_url=f"https://github.com/acme/{slug}",
        summary="Operational catalog fixture.",
        compatible_runtimes=["claude", "codex"],
        metadata={
            "github_stars": stars,
            "source_accessible": True,
            "detection_succeeded": True,
            "skill_paths": ["skills/demo/SKILL.md"] if skill else [],
        },
    )


def _web_resource(*, slug: str) -> ResourceItem:
    return ResourceItem(
        slug=slug,
        name=slug.replace("-", " ").title(),
        kind=ResourceKind.REFERENCE,
        source="web",
        source_url=f"https://example.com/{slug}",
        summary="Operational web reference.",
        compatible_runtimes=["claude", "codex"],
        metadata={"source_accessible": True, "detection_succeeded": True},
    )


def test_default_policy_is_strict_and_accepted_only_filters(karajan_install, monkeypatch, tmp_path) -> None:
    client, store = _client(karajan_install, monkeypatch, tmp_path)
    accepted = store.upsert(_resource(slug="accepted", stars=10_001, skill=True))
    pending = store.upsert(_resource(slug="pending", stars=10_000))

    policy = client.get("/resources/validity-policy")
    assert policy.status_code == 200
    assert policy.json()["min_github_stars"] == 10_000
    assert policy.json()["required_runtimes"] == ["claude", "codex"]
    visible = client.get("/resources", params={"accepted_only": True}).json()
    assert [item["id"] for item in visible] == [accepted.id]
    review = client.get("/resources/review").json()
    assert [item["id"] for item in review] == [pending.id]
    assert review[0]["metadata"]["validity"]["failed_checks"] == ["stars"]


def test_web_references_are_exempt_from_github_stars(karajan_install, monkeypatch, tmp_path) -> None:
    client, store = _client(karajan_install, monkeypatch, tmp_path)
    web = store.upsert(_web_resource(slug="web-reference"))

    visible = client.get("/resources", params={"accepted_only": True}).json()
    assert [item["id"] for item in visible] == [web.id]
    validity = visible[0]["metadata"]["validity"]
    assert validity["github_repository"] is False
    assert validity["github_stars"] is None
    assert validity["checks"]["stars"]["enabled"] is False
    assert validity["checks"]["stars"]["passed"] is True


def test_unknown_github_stars_remain_pending_not_zero(karajan_install, monkeypatch, tmp_path) -> None:
    client, store = _client(karajan_install, monkeypatch, tmp_path)
    item = _resource(slug="unknown-stars", stars=1).model_copy(
        update={"metadata": {"source_accessible": True, "detection_succeeded": True}}
    )
    stored = store.upsert(item)

    review = client.get("/resources/review").json()
    assert [entry["id"] for entry in review] == [stored.id]
    validity = review[0]["metadata"]["validity"]
    assert validity["github_stars"] is None
    assert validity["checks"]["stars"]["actual"] is None
    assert validity["failed_checks"] == ["stars"]


def test_internal_repository_components_use_parent_state_and_are_not_review_items(tmp_path) -> None:
    store = ResourceStore(tmp_path / "components.db")
    parent = store.upsert(_resource(slug="ecc", stars=244_000, skill=True))
    child = store.upsert(ResourceItem(
        slug="ecc--agent-introspection-debugging", name="Agent introspection debugging",
        kind=ResourceKind.SKILL, source="github-skill", parent_resource_id=parent.id,
        source_url="https://github.com/affaan-m/ECC/blob/main/docs/ja-JP/skills/agent-introspection-debugging/SKILL.md",
        compatible_runtimes=["claude", "codex"], metadata={
            "skill_paths": ["docs/ja-JP/skills/agent-introspection-debugging/SKILL.md"],
            "provenance": {"method": "github-tree-skill-path"},
        },
    ))
    decorated = {item.id: item for item in decorate_catalog_resources(store.list_all_resources(), load_policy(store))}
    assert decorated[parent.id].metadata["validity"]["state"] == "accepted"
    assert decorated[child.id].metadata["validity"]["state"] == "accepted"
    assert decorated[child.id].metadata["validity"]["github_stars"] == 244_000
    assert decorated[child.id].metadata["validity"]["analysis_owner_id"] == parent.id


def test_local_skill_alias_uses_canonical_github_resource(tmp_path) -> None:
    store = ResourceStore(tmp_path / "aliases.db")
    canonical = store.upsert(_resource(slug="aider", stars=48_531, skill=False).model_copy(update={"name": "aider", "source": "naityer"}))
    local = store.upsert(ResourceItem(
        slug="aider-local", name="aider", kind=ResourceKind.SKILL, source="skills_catalog",
        summary="Local copy", compatible_runtimes=["claude", "codex"], metadata={
            "source_analysis_method": "local-files", "source_revision": "local",
        },
    ))
    decorated = {item.id: item for item in decorate_catalog_resources(store.list_all_resources(), load_policy(store))}
    assert decorated[local.id].metadata["validity"]["alias_of"] == canonical.id
    assert decorated[local.id].metadata["validity"]["state"] == "accepted"
    assert decorated[local.id].metadata["validity"]["github_stars"] == 48_531
    assert decorated[local.id].metadata["validity"]["source_url"] == canonical.repo_url


def test_nested_web_pages_use_their_directory_parent(tmp_path) -> None:
    store = ResourceStore(tmp_path / "web-components.db")
    parent = store.upsert(ResourceItem(
        slug="refactoring-guru", name="Refactoring Guru", kind=ResourceKind.REFERENCE,
        source="refactoring.guru", source_url="https://refactoring.guru/design-patterns",
        compatible_runtimes=["claude", "codex"], metadata={
            "source_accessible": True, "detection_succeeded": True,
        },
    ))
    child = store.upsert(ResourceItem(
        slug="refactoring-guru--chain-of-responsibility", name="Chain of Responsibility",
        kind=ResourceKind.REFERENCE, source="refactoring.guru", parent_resource_id=parent.id,
        source_url="https://refactoring.guru/design-patterns/chain-of-responsibility",
        compatible_runtimes=[], metadata={
            "provenance": {"method": "generic-html", "parent_resource_id": parent.id},
        },
    ))
    decorated = {item.id: item for item in decorate_catalog_resources(store.list_all_resources(), load_policy(store))}
    assert decorated[parent.id].metadata["validity"]["state"] == "accepted"
    assert decorated[child.id].metadata["validity"]["state"] == "accepted"
    assert decorated[child.id].metadata["validity"]["analysis_owner_id"] == parent.id


def test_verified_curated_web_source_counts_as_accessible(tmp_path) -> None:
    item = ResourceItem(
        slug="curated-web", name="Curated web", kind=ResourceKind.REFERENCE,
        source="curated", source_url="https://refactoring.guru/design-patterns/catalog",
        compatible_runtimes=["claude", "codex"], metadata={
            "verified_source": True, "detection_succeeded": True,
        },
    )
    validity = decorate_catalog_resources([item], load_policy(ResourceStore(tmp_path / "verified.db")))[0].metadata["validity"]
    assert validity["checks"]["source"]["passed"] is True


def test_curated_web_directory_infers_supported_runtimes(tmp_path) -> None:
    item = ResourceItem(
        slug="curated-directory", name="Curated directory", kind=ResourceKind.REFERENCE,
        source="curated", source_url="https://refactoring.guru/design-patterns/catalog",
        metadata={"verified_source": True, "source_role": "collection", "source_revision": "sha"},
    )
    validity = decorate_catalog_resources([item], load_policy(ResourceStore(tmp_path / "runtimes.db")))[0].metadata["validity"]
    assert validity["checks"]["runtimes"]["actual"] == ["claude", "codex"]
    assert validity["checks"]["runtimes"]["passed"] is True


def test_impact_is_read_only_and_policy_save_resets_decisions(karajan_install, monkeypatch, tmp_path) -> None:
    client, store = _client(karajan_install, monkeypatch, tmp_path)
    item = store.upsert(_resource(slug="manual", stars=2))
    accepted = client.post(f"/resources/{item.id}/review", json={"decision": "accept"})
    assert accepted.status_code == 200
    before = store.get(item.id).metadata
    draft = {
        "stars_enabled": False,
        "min_github_stars": 10_000,
        "require_source_url": True,
        "require_detection_success": True,
        "required_runtimes": ["claude", "codex"],
        "require_web_reference": True,
    }
    impact = client.post("/resources/validity-policy/impact", json=draft)
    assert impact.status_code == 200
    assert impact.json()["manual_decisions_to_reset"] == 1
    assert store.get(item.id).metadata == before

    denied = client.put("/resources/validity-policy", json={**draft, "expected_version": 1, "confirm_reset": False})
    assert denied.status_code == 409
    saved = client.put("/resources/validity-policy", json={**draft, "expected_version": 1, "confirm_reset": True})
    assert saved.status_code == 200
    assert saved.json()["policy"]["version"] == 2
    validity = store.get(item.id).metadata["validity"]
    assert validity["decision"] == "automatic"
    assert validity["history"][-1]["reason"] == "policy_changed"
    conflict = client.put("/resources/validity-policy", json={**draft, "expected_version": 1, "confirm_reset": True})
    assert conflict.status_code == 409


def test_review_decisions_and_excel_snapshot_are_auditable(karajan_install, monkeypatch, tmp_path) -> None:
    client, store = _client(karajan_install, monkeypatch, tmp_path)
    accepted = store.upsert(_resource(slug="xlsx-accepted", stars=20_000))
    rejected = store.upsert(_resource(slug="xlsx-review", stars=9))
    assert client.post(f"/resources/{rejected.id}/review", json={"decision": "discard"}).json()["metadata"]["validity"]["state"] == "discarded"

    preview = client.post("/resources/inventory")
    assert preview.status_code == 200
    body = preview.json()
    assert body["total"] == 2
    assert len(body["rows"]) == 2
    download = client.get(body["download_url"])
    assert download.status_code == 200
    workbook = load_workbook(BytesIO(download.content))
    assert workbook.sheetnames == ["Inventario", "Política"]
    inventory = workbook["Inventario"]
    assert inventory.freeze_panes == "A2"
    names = {inventory.cell(row=row, column=1).value for row in range(2, inventory.max_row + 1)}
    assert names == {accepted.name, rejected.name}
    states = {inventory.cell(row=row, column=10).value for row in range(2, inventory.max_row + 1)}
    assert states == {"accepted", "discarded"}


def test_claude_collection_keeps_original_links_and_deduplicates() -> None:
    markdown = b"""# Collection
| Skill | Source |
| --- | --- |
| Diagram Design | [Source](https://github.com/cathrynlavery/diagram-design) |
| Diagram duplicate | [Repo](https://github.com/cathrynlavery/diagram-design/) |
| Nested Skill | [View](https://github.com/acme/tools/tree/main/skills/nested) |
"""

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, content=markdown, request=request)

    parent = ResourceItem(
        slug="claude-skills-collection", name="Claude Skills Collection",
        kind=ResourceKind.REFERENCE, source="curated",
        repo_url="https://github.com/abubakarsiddik31/claude-skills-collection",
    )
    children = ClaudeSkillsCollectionAdapter(httpx.Client(transport=httpx.MockTransport(handler))).discover(parent)
    assert len(children) == 2
    assert {child.source_url for child in children} == {
        "https://github.com/cathrynlavery/diagram-design",
        "https://github.com/acme/tools/tree/main/skills/nested",
    }
    assert all(child.parent_resource_id == parent.id for child in children)
    assert all(child.metadata["provenance"]["collection_url"].endswith("claude-skills-collection") for child in children)
